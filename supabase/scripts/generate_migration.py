#!/usr/bin/env python3
"""Genera i file SQL di migrazione dal foglio Google esportato.

Uso:
    python3 supabase/scripts/generate_migration.py "new_structure/SOCI 2026.xlsx"
    python3 supabase/scripts/generate_migration.py "SOCI 2026.xlsx" supabase/migration 300

Argomenti: file .xlsx di partenza, cartella di destinazione (default
supabase/migration), righe per file (default 500 — abbassarlo se l'editor SQL
di Supabase rifiuta ancora il file).

Produce più file numerati da eseguire in ordine, invece di un unico file da
1,2 MB che l'editor SQL rifiuta ("Query is too large"). Ogni file ha la propria
transazione ed è rieseguibile senza creare duplicati.

Richiede openpyxl. Rigenerare dopo ogni nuovo export del foglio.
"""

import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

# Colonne del foglio SOCI 1.x (1-indexed), come lette da codice.gs/Admin.gs.
COL = {
    "timestamp": 1,
    "numero_polizza": 2,
    "cognome": 3,
    "nome": 4,
    "luogo_nascita": 5,
    "provincia_nascita": 6,
    "codice_fiscale": 7,
    "data_nascita": 8,
    "indirizzo": 9,
    "citta": 10,
    "provincia": 11,
    "cap": 12,
    "telefono": 13,
    "email": 14,
    "tipologia_tessera": 15,
    "agevolazioni_famiglia": 16,
    "tipo_abbonamento": 17,
    "partenza_domenica": 18,
    "partenze_sabato": 19,
    "tipologia_corso": 20,
    # Attenzione: nel foglio le etichette di colonna 21 e 23 sono invertite
    # rispetto al contenuto. Qui si segue il contenuto reale, cioè quello che
    # codice.gs e Admin.gs scrivono davvero.
    "totale": 21,
    "acconto": 22,
    "saldo_ignorato": 23,  # ricalcolato come colonna generata, non migrato
    "legacy_id": 24,
    "numero_tessera": 25,
    "legacy_payer_id": 26,
}

# Anagrafiche per file. 500 righe stanno intorno ai 140 KB, sotto il limite
# dell'editor SQL di Supabase. Si può abbassare da riga di comando.
RIGHE_PER_FILE = 500

# Le righe importate sono un archivio anagrafico: nel foglio la colonna
# "Informazioni cronologiche" è vuota al 100%, quindi non esiste una data di
# iscrizione reale da conservare. Se si lasciasse il default (now()) tutte le
# 5.000+ anagrafiche storiche risulterebbero iscritte alla stagione corrente e
# falserebbero il riepilogo. Si assegna quindi una data d'archivio esplicita,
# precedente a qualsiasi stagione: le nuove iscrizioni avranno la data vera.
DATA_ARCHIVIO = "2000-01-01T00:00:00+00:00"

TEXT_FIELDS = [
    "numero_polizza", "cognome", "nome", "luogo_nascita", "provincia_nascita",
    "codice_fiscale", "indirizzo", "citta", "provincia", "cap", "telefono",
    "email", "tipologia_tessera", "agevolazioni_famiglia", "tipo_abbonamento",
    "partenza_domenica", "partenze_sabato", "tipologia_corso", "numero_tessera",
]


def sql_str(value):
    """Letterale SQL per un testo, o NULL."""
    if value is None:
        return "NULL"
    text = str(value).strip()
    if text == "":
        return "NULL"
    return "'" + text.replace("'", "''") + "'"


def sql_num(value):
    if value is None or value == "":
        return "0"
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return "0"


def clean_phone(value):
    """I telefoni sono letti da Excel come float (3468724757.0)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_date(value):
    """Replica il parser di getMembers() in codice.gs.

    Formati presenti nel foglio: datetime nativo, 'mm/dd/yyyy', 'mm\\dd\\yy'.
    L'anno a due cifre segue la stessa regola della 1.x: <30 -> 2000+, <100 -> 1900+.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip().replace("\\", "/")
    parts = [p for p in re.split(r"[/\-.]", text) if p != ""]
    if len(parts) != 3:
        return None
    try:
        month, day, year = int(parts[0]), int(parts[1]), int(parts[2])
    except ValueError:
        return None

    if year < 30:
        year += 2000
    elif year < 100:
        year += 1900

    # Alcune righe hanno giorno e mese invertiti: se il mese non è valido ma il
    # giorno sì, si scambiano invece di scartare la data.
    if month > 12 and day <= 12:
        month, day = day, month
    try:
        return date(year, month, day)
    except ValueError:
        return None


def read_soci(worksheet):
    rows = []
    for row_index in range(2, worksheet.max_row + 1):
        cognome = worksheet.cell(row=row_index, column=COL["cognome"]).value
        nome = worksheet.cell(row=row_index, column=COL["nome"]).value
        if not cognome and not nome:
            continue

        def cell(field):
            return worksheet.cell(row=row_index, column=COL[field]).value

        record = {field: cell(field) for field in TEXT_FIELDS}
        record["telefono"] = clean_phone(record["telefono"])
        record["data_nascita"] = parse_date(cell("data_nascita"))
        record["totale"] = cell("totale")
        record["acconto"] = cell("acconto")
        record["legacy_id"] = cell("legacy_id")
        record["legacy_payer_id"] = cell("legacy_payer_id")
        record["timestamp"] = cell("timestamp")
        record["_row"] = row_index
        rows.append(record)
    return rows


def read_prezzi(worksheet):
    out = []
    for row_index in range(2, worksheet.max_row + 1):
        categoria = worksheet.cell(row=row_index, column=1).value
        nome = worksheet.cell(row=row_index, column=2).value
        prezzo = worksheet.cell(row=row_index, column=3).value
        if not categoria or not nome:
            continue
        out.append((str(categoria).strip().upper(), str(nome).strip(), prezzo))
    return out


def read_partenze(worksheet):
    out = []
    for row_index in range(2, worksheet.max_row + 1):
        giorno = worksheet.cell(row=row_index, column=1).value
        luogo = worksheet.cell(row=row_index, column=2).value
        if not giorno or not luogo:
            continue
        out.append((str(giorno).strip().upper(), str(luogo).strip()))
    return out


def intestazione(titolo, numero, totale, source_name, extra=()):
    """Intestazione comune a ogni file: dice cos'è e in che ordine va eseguito."""
    righe = [
        f"-- Sci Club Don Bosco 2.0 — {titolo}",
        f"-- File {numero} di {totale} — eseguire i file in ordine di numero.",
        f"-- Generato da {source_name} con supabase/scripts/generate_migration.py",
        "--",
        "-- Prerequisito: supabase/schema.sql già eseguito.",
        "-- Ogni file è indipendente e rieseguibile: rilanciarlo non crea duplicati.",
    ]
    righe += [f"-- {r}" for r in extra]
    righe += ["", "begin;", ""]
    return righe


def riga_socio(record, columns_count=None):
    """Una tupla VALUES per un socio."""
    data_nascita = record["data_nascita"]
    values = ", ".join([
        f"TIMESTAMPTZ '{DATA_ARCHIVIO}'",
        f"TIMESTAMPTZ '{DATA_ARCHIVIO}'",
        sql_str(record["legacy_id"]),
        sql_str(record["legacy_payer_id"]),
        sql_str(record["numero_polizza"]),
        sql_str(record["cognome"]),
        sql_str(record["nome"]),
        sql_str(record["luogo_nascita"]),
        sql_str(record["provincia_nascita"]),
        sql_str(record["codice_fiscale"]),
        f"DATE '{data_nascita.isoformat()}'" if data_nascita else "NULL",
        sql_str(record["indirizzo"]),
        sql_str(record["citta"]),
        sql_str(record["provincia"]),
        sql_str(record["cap"]),
        sql_str(record["telefono"]),
        sql_str(record["email"]),
        sql_str(record["tipologia_tessera"]),
        sql_str(record["agevolazioni_famiglia"]),
        sql_str(record["tipo_abbonamento"]),
        sql_str(record["partenza_domenica"]),
        sql_str(record["partenze_sabato"]),
        sql_str(record["tipologia_corso"]),
        sql_num(record["totale"]),
        sql_num(record["acconto"]),
        sql_str(record["numero_tessera"]),
    ])
    return f"  ({values})"


def build_files(soci, prezzi, partenze, source_name, righe_per_file):
    """Produce la lista (nome_file, contenuto) dei file di migrazione.

    L'editor SQL di Supabase rifiuta le query troppo grandi, quindi i soci
    vengono divisi in più file. Ogni file apre e chiude la propria transazione:
    si possono lanciare uno alla volta, anche a distanza di tempo, e rilanciare
    senza produrre duplicati.
    """
    blocchi_soci = [soci[i:i + righe_per_file] for i in range(0, len(soci), righe_per_file)]
    # setup + blocchi soci + finalizzazione
    totale_file = len(blocchi_soci) + 2
    files = []

    # ---------------------------------------------------------------- setup
    lines = intestazione(
        "listino, partenze e preparazione", 1, totale_file, source_name,
        extra=[
            "Carica il listino prezzi e i luoghi di partenza.",
        ],
    )
    lines += ["-- Listino", ""]
    for categoria, nome, prezzo in prezzi:
        lines.append(
            "insert into public.prices (category, name, price) values "
            f"({sql_str(categoria)}, {sql_str(nome)}, {sql_num(prezzo)}) "
            "on conflict (category, name) do update set price = excluded.price;"
        )

    lines += ["", "-- Luoghi di partenza", ""]
    for giorno, luogo in partenze:
        lines.append(
            "insert into public.departures (day, place) values "
            f"({sql_str(giorno)}, {sql_str(luogo)}) on conflict (day, place) do nothing;"
        )

    lines += [
        "",
        "commit;",
        "",
    ]
    files.append(("01_listino_e_partenze.sql", "\n".join(lines)))

    # ----------------------------------------------------------- soci
    columns = (
        "created_at, enrolled_at, legacy_id, legacy_payer_id, policy_number, last_name, first_name, birth_place, "
        "birth_province, tax_code, birth_date, address, city, province, "
        "postal_code, phone, email, card_type, family_discount, pass_type, "
        "sunday_departure, saturday_departure, course_type, total, paid, card_number"
    )

    for indice, blocco in enumerate(blocchi_soci):
        numero_file = indice + 2
        primo = indice * righe_per_file + 1
        ultimo = primo + len(blocco) - 1

        lines = intestazione(
            f"soci {primo}-{ultimo}", numero_file, totale_file, source_name,
            extra=[
                f"{len(blocco)} anagrafiche su {len(soci)} totali.",
                "I collegamenti familiari NON vengono risolti qui: lo fa l'ultimo file.",
            ],
        )
        lines.append(f"insert into public.members ({columns}) values")
        lines.append(",\n".join(riga_socio(r) for r in blocco))
        lines.append("on conflict (legacy_id) do nothing;")
        lines += ["", "commit;", ""]

        files.append((f"{numero_file:02d}_soci_{primo:05d}_{ultimo:05d}.sql", "\n".join(lines)))

    # -------------------------------------------------------- finalizzazione
    lines = intestazione(
        "collegamenti familiari", totale_file, totale_file, source_name,
        extra=[
            "Da eseguire SOLO dopo tutti i file dei soci: collega ogni socio al",
            "proprio capofamiglia.",
        ],
    )
    lines += [
        "-- Risoluzione dei collegamenti: legacy_payer_id -> payer_id",
        "update public.members d",
        "   set payer_id = c.id",
        "  from public.members c",
        " where d.legacy_payer_id is not null",
        "   and d.legacy_payer_id <> ''",
        "   and c.legacy_id = d.legacy_payer_id",
        "   and c.id <> d.id;",
        "",
        "-- Collegamenti rimasti irrisolti (pagante non presente nel foglio):",
        "-- vengono segnalati e lasciati senza payer_id, non scartati.",
        "do $$",
        "declare orfani int;",
        "begin",
        "  select count(*) into orfani",
        "    from public.members",
        "   where legacy_payer_id is not null and legacy_payer_id <> '' and payer_id is null;",
        "  if orfani > 0 then",
        "    raise notice 'Collegamenti familiari non risolti: %', orfani;",
        "  end if;",
        "end $$;",
        "",
        "commit;",
        "",
        "-- Controllo finale: quante anagrafiche sono state caricate.",
        "select count(*) as soci_caricati,",
        "       count(payer_id) as con_capofamiglia",
        "  from public.members;",
        "",
    ]
    files.append((f"{totale_file:02d}_collegamenti_familiari.sql", "\n".join(lines)))

    return files


def main():
    if not 2 <= len(sys.argv) <= 4:
        print(__doc__)
        return 1

    source = Path(sys.argv[1])
    target_dir = Path(sys.argv[2] if len(sys.argv) > 2 else "supabase/migration")
    righe_per_file = int(sys.argv[3]) if len(sys.argv) > 3 else RIGHE_PER_FILE

    workbook = openpyxl.load_workbook(source, data_only=True)
    soci = read_soci(workbook["SOCI"])
    prezzi = read_prezzi(workbook["PREZZI"])
    partenze = read_partenze(workbook["PARTENZE"])

    target_dir.mkdir(parents=True, exist_ok=True)
    # Rimuove i file di una generazione precedente: se il numero di blocchi
    # cala, i vecchi file resterebbero lì e verrebbero eseguiti per sbaglio.
    for vecchio in target_dir.glob("*.sql"):
        vecchio.unlink()

    files = build_files(soci, prezzi, partenze, source.name, righe_per_file)
    for nome, contenuto in files:
        (target_dir / nome).write_text(contenuto, encoding="utf-8")

    senza_data = sum(1 for r in soci if r["data_nascita"] is None)
    print(f"soci: {len(soci)} (senza data di nascita valida: {senza_data})")
    print(f"prezzi: {len(prezzi)}  partenze: {len(partenze)}")
    print(f"\n{len(files)} file scritti in {target_dir}/ — eseguirli in ordine:\n")
    for nome, contenuto in files:
        kb = len(contenuto.encode("utf-8")) / 1024
        print(f"  {nome:40s} {kb:7.0f} KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
